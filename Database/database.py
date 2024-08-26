import pandas as pd
import mysql.connector
from mysql.connector import Error
import math
import warnings
import numpy as np
from openpyxl import load_workbook
import psycopg2
from psycopg2 import OperationalError
import logging

class DB:
    def __init__(self, user=None, password=None, host=None, database_name=None):
        self.user = user
        self.password = password
        self.host = host
        self.database = database_name
        self.mycursor = None
        self.mydatabase = None

    def connectDb(self):
        """
        Establishes a connection to a MySQL database and sets up a cursor for executing queries.

        :return: None

        Owner: Prashant Sahu
        Date: 2024-08-26
        """
        try:
            # Attempt to establish a connection to the MySQL database
            self.mydatabase = mysql.connector.connect(
                user=self.user,
                password=self.password,
                host=self.host,
                database=self.database
            )
            self.mycursor = self.mydatabase.cursor()

            # Check if the connection was successful
            if self.mydatabase.is_connected():
                print("Connected to the database successfully!")
            else:
                raise ConnectionError("Failed to establish a database connection.")

        except mysql.connector.Error as db_error:
            # Handle specific database connection errors
            logging.error(f"Database connection error: {db_error}")
            self._cleanup_resources()

        except ConnectionError as conn_error:
            # Handle general connection errors
            logging.error(f"Connection error: {conn_error}")
            self._cleanup_resources()

        except Exception as e:
            # Handle any other unexpected errors
            logging.error(f"An unexpected error occurred: {e}")
            self._cleanup_resources()

    def connectDbPostgres(self):
        """
        Establishes a connection to a PostgreSQL database and sets up a cursor for executing queries.

        :return: None

        Owner: Prashant Sahu
        Date: 2024-08-26
        """
        try:
            # Attempt to establish a connection to the PostgreSQL database
            self.mydatabase = psycopg2.connect(
                user=self.user,
                password=self.password,
                host=self.host,
                database=self.database
            )
            self.mycursor = self.mydatabase.cursor()

            # Check if the connection was successful
            if self.mydatabase:
                print("Connected to the database successfully!")
            else:
                raise ConnectionError("Failed to establish a database connection.")

        except OperationalError as db_error:
            # Handle errors specific to database operations
            logging.error(f"OperationalError during PostgreSQL connection: {db_error}")
            self._cleanup_resources()

        except ConnectionError as conn_error:
            # Handle general connection errors
            logging.error(f"Connection error: {conn_error}")
            self._cleanup_resources()

        except Exception as e:
            # Handle any other unexpected errors
            logging.error(f"An unexpected error occurred: {e}")
            self._cleanup_resources()

    def is_nan(self, x):
        return isinstance(x, float) and math.isnan(x)

    def insertDataToDb(self, df, table_name, create_table_query):
        """
        Inserts data from a pandas DataFrame into a MySQL database table.
        Owner: Prashant Sahu
        Date: 2024-08-26
        """
        try:
            # Check if DataFrame is empty
            if len(df) == 0:
                print("DataFrame is empty. Nothing to insert.")
                return

            # Drop the table if it already exists and create a new one
            self.mycursor.execute(f"DROP TABLE IF EXISTS {table_name}")
            self.mycursor.execute(create_table_query)

            # Replace NaN values with None for SQL compatibility
            df = df.where(pd.notnull(df), None)

            # Get column names from DataFrame and prepare SQL insert statement
            columns = df.columns.tolist()
            columns_str = ", ".join(columns)
            placeholders = ", ".join(["%s"] * len(columns))
            sql = f"INSERT INTO {table_name} ({columns_str}) VALUES ({placeholders})"

            # Insert data row by row
            for _, row in df.iterrows():
                values = [row[col] if pd.notnull(row[col]) else None for col in columns]
                self.mycursor.execute(sql, values)

            # Commit the transaction to the database
            self.mydatabase.commit()
            print("Data inserted successfully.")
        except Exception as e:
            print(f"Error inserting data to DB: {e}")

    def readDatabase(self, table_name, query=None) -> pd.DataFrame:
        """
        Reads data from a MySQL database into a pandas DataFrame.

        :param table_name: Name of the table to read data from.
        :param query: Optional SQL query string. If provided, the query will be executed
                      instead of selecting all rows from the table.
        :return: A pandas DataFrame containing the data read from the database.

        Owner: Prashant Sahu
        Date: 2024-08-26
        """
        try:
            df = None
            # Suppress the specific UserWarning from pandas
            with warnings.catch_warnings():
                warnings.simplefilter("ignore", UserWarning)

                # Read data from MySQL into a DataFrame using the provided query or table name
                if query is not None:
                    df = pd.read_sql(query, con=self.mydatabase)
                else:
                    df = pd.read_sql(f"SELECT * FROM {table_name}", con=self.mydatabase)

                # Replace None values with NaN in the DataFrame
                df.replace(to_replace=[None], value=float('nan'), inplace=True)

            return df
        except Exception as e:
            print(f"Error reading source data from DB: {e}")
            return df

    def closeDb(self):
        try:
            if self.mycursor:
                self.mycursor.close()
            if self.mydatabase:
                self.mydatabase.close()
            print("Database connection closed successfully!")
        except Exception as e:
            print(f"Error closing the database connection: {e}")

    def _cleanup_resources(self):
        # Clean up resources if there was an error
        if self.mycursor is not None:
            self.mycursor.close()
            self.mycursor = None
        if self.mydatabase is not None and self.mydatabase.is_connected():
            self.mydatabase.close()
            self.mydatabase = None







    



    
    # def compare_tables(self):
        # try:
        #     target_tuples = list(self.target_df.to_records(index=False)) 
        #     source_tuples = list(self.source_df.to_records(index=False))

        #     differce_count = 0
        #     min_length = min(len(source_tuples), len(target_tuples))
        #     for i in range(min_length):
        #         target_row = target_tuples[i]
        #         source_row = source_tuples[i]

        #         for j in range(len(target_row)):
        #             if (self.is_nan(target_row[j]) and self.is_nan(source_row[j])):
        #                 continue
        #             elif target_row[j] != source_row[j]:
        #                 differce_count += 1
        #                 break

        #     differce_count += max(len(source_tuples), len(target_tuples)) - min_length

        #     print(f"Number of differing rows: {differce_count}")
        # except Exception as e:
        #     print(f"Error during table comparison: {e}")
   
    